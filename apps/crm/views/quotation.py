"""
Quotation views: CRUD with RBAC enforcement.
"""
from django.contrib import messages
from django.contrib.auth.mixins import LoginRequiredMixin
from django.views.generic import ListView, DetailView, CreateView, UpdateView, DeleteView
from django.urls import reverse_lazy

from apps.core.mixins import OwnerRequiredMixin, OwnerOrAdminMixin, SetOwnerMixin, AdminOnlyMixin
from apps.crm.models import Quotation
from apps.crm.forms import QuotationForm


class QuotationListView(OwnerRequiredMixin, ListView):
    model = Quotation
    template_name = 'crm/quotations/quotation_list.html'
    context_object_name = 'quotations'
    paginate_by = 15

    def get_queryset(self):
        qs = super().get_queryset().select_related('lead__client', 'owner')
        search = self.request.GET.get('q', '').strip()
        start_date = self.request.GET.get('start_date', '').strip()
        end_date = self.request.GET.get('end_date', '').strip()
        company = self.request.GET.get('company', '').strip()
        status = self.request.GET.get('status', '').strip()
        
        from apps.core.date_utils import parse_date_safe
        psd = parse_date_safe(start_date)
        ped = parse_date_safe(end_date)
        
        # Exclude DEVELOPMENT testing data
        qs = qs.exclude(lead__client__company_name__iexact='DEVELOPMENT')

        if search:
            search_upper = search.upper()
            if 'QUO-' in search_upper and search_upper.split('QUO-')[1].isdigit():
                matched_qs = qs.filter(pk=int(search_upper.split('QUO-')[1]))
            else:
                from django.db.models import Q
                matched_qs = qs.filter(
                    Q(lead__client__company_name__icontains=search) |
                    Q(lead__client__contact_person__icontains=search) |
                    Q(origin__icontains=search) |
                    Q(destination__icontains=search)
                )
            filtered_qs = matched_qs
            if status:
                filtered_qs = filtered_qs.filter(status=status)
            if company:
                filtered_qs = filtered_qs.filter(lead__client__company_name__icontains=company)
                
            date_filtered_qs = filtered_qs
            if psd:
                date_filtered_qs = date_filtered_qs.filter(created_at__date__gte=psd)
            if ped:
                date_filtered_qs = date_filtered_qs.filter(created_at__date__lte=ped)
                
            if date_filtered_qs.exists():
                return date_filtered_qs
            elif filtered_qs.exists():
                return filtered_qs
            return matched_qs
            
        if status:
            qs = qs.filter(status=status)
        if psd:
            qs = qs.filter(created_at__date__gte=psd)
        if ped:
            qs = qs.filter(created_at__date__lte=ped)
        if company:
            qs = qs.filter(lead__client__company_name__icontains=company)
            
        return qs

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['search_query'] = self.request.GET.get('q', '')
        context['start_date'] = self.request.GET.get('start_date', '')
        context['end_date'] = self.request.GET.get('end_date', '')
        context['company_query'] = self.request.GET.get('company', '')
        context['status_filter'] = self.request.GET.get('status', '')
        context['status_choices'] = Quotation.Status.choices
        from apps.crm.models import Client
        context['available_companies'] = Client.objects.exclude(company_name='DEVELOPMENT').values_list('company_name', flat=True).distinct()
        return context


class QuotationDetailView(OwnerOrAdminMixin, DetailView):
    model = Quotation
    template_name = 'crm/quotations/quotation_detail.html'
    context_object_name = 'quotation'


from apps.crm.forms import QuotationForm, QuotationItemFormSet

class QuotationCreateView(LoginRequiredMixin, SetOwnerMixin, CreateView):
    model = Quotation
    form_class = QuotationForm
    template_name = 'crm/quotations/quotation_edit.html'

    def get_success_url(self):
        return self.object.lead.get_absolute_url()

    def get_initial(self):
        initial = super().get_initial()
        lead_id = self.request.GET.get('lead')
        if lead_id:
            initial['lead'] = lead_id
        return initial

    def get_form_kwargs(self):
        kwargs = super().get_form_kwargs()
        kwargs['user'] = self.request.user
        return kwargs
        
    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        if self.request.POST:
            context['items_formset'] = QuotationItemFormSet(self.request.POST)
        else:
            context['items_formset'] = QuotationItemFormSet()
        return context

    def form_valid(self, form):
        context = self.get_context_data()
        items_formset = context['items_formset']
        
        if items_formset.is_valid():
            # Allow SetOwnerMixin to set the owner and save the Quotation instance
            response = super().form_valid(form)
            
            # Now self.object is safely saved, we can save the items
            items_formset.instance = self.object
            items_formset.save()
            
            # Recalculate totals based on items
            self.object.recalculate_totals()
            
            messages.success(self.request, 'Penawaran berhasil dibuat.')
            return response
        else:
            return self.form_invalid(form)


class QuotationUpdateView(OwnerOrAdminMixin, UpdateView):
    model = Quotation
    form_class = QuotationForm
    template_name = 'crm/quotations/quotation_edit.html'
    success_url = reverse_lazy('crm:quotation-list')

    def get_form_kwargs(self):
        kwargs = super().get_form_kwargs()
        kwargs['user'] = self.request.user
        return kwargs

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        if self.request.POST:
            context['items_formset'] = QuotationItemFormSet(self.request.POST, instance=self.object)
        else:
            context['items_formset'] = QuotationItemFormSet(instance=self.object)
        return context

    def form_valid(self, form):
        context = self.get_context_data()
        items_formset = context['items_formset']
        
        if items_formset.is_valid():
            # Save the main form first
            response = super().form_valid(form)
            
            # Save the items formset
            items_formset.instance = self.object
            items_formset.save()
            
            # Recalculate totals
            self.object.recalculate_totals()
            
            messages.success(self.request, 'Quotation berhasil diperbarui.')
            return response
        else:
            return self.form_invalid(form)


class QuotationDeleteView(AdminOnlyMixin, DeleteView):
    model = Quotation
    template_name = 'crm/quotations/quotation_delete.html'
    success_url = reverse_lazy('crm:quotation-list')

    def form_valid(self, form):
        messages.success(self.request, 'Quotation berhasil dihapus.')
        return super().form_valid(form)


class QuotationPrintView(OwnerOrAdminMixin, DetailView):
    model = Quotation
    template_name = 'crm/quotations/quotation_print.html'
    context_object_name = 'quotation'


class QuotationExportExcelView(OwnerRequiredMixin, ListView):
    """Export Quotations to Excel (.xlsx) with styling."""
    model = Quotation

    def get(self, request, *args, **kwargs):
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from django.http import HttpResponse
        from django.utils import timezone as tz

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Quotations Export"

        title_font = Font(size=16, bold=True, color="000000")
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="343A40", end_color="343A40", fill_type="solid")
        border = Border(
            left=Side(style='thin'), right=Side(style='thin'),
            top=Side(style='thin'), bottom=Side(style='thin')
        )
        align_center = Alignment(horizontal="center", vertical="center")
        align_left = Alignment(horizontal="left", vertical="center")

        ws.merge_cells('A1:K1')
        ws['A1'].value = "DATA PENAWARAN (QUOTATION)"
        ws['A1'].font = title_font
        ws['A1'].alignment = align_center

        ws.merge_cells('A2:K2')
        ws['A2'].value = f"Diekspor pada: {tz.now().strftime('%d %b %Y')}"
        ws['A2'].alignment = align_center
        ws.append([])

        headers = [
            'No. Quotation', 'Perusahaan', 'Layanan', 'Asal', 'Tujuan',
            'Berat Aktual (kg)', 'Berat Vol. (kg)', 'Harga/Kg (Rp)',
            'Total (Rp)', 'PIC', 'Tanggal'
        ]
        ws.append(headers)
        for cell in ws[4]:
            cell.font = header_font
            cell.fill = header_fill
            cell.border = border
            cell.alignment = align_center

        qs = self.get_queryset().select_related('lead__client', 'owner')
        search = request.GET.get('q', '')
        if search:
            from django.db.models import Q
            qs = qs.filter(
                Q(lead__client__company_name__icontains=search) |
                Q(origin__icontains=search) |
                Q(destination__icontains=search)
            )

        for row_num, q in enumerate(qs, start=5):
            row_data = [
                q.quotation_number,
                q.lead.client.company_name,
                q.get_service_display() if q.service else '-',
                q.origin or '-',
                q.destination or '-',
                float(q.actual_weight),
                float(q.volumetric_weight),
                float(q.price_per_kg),
                float(q.total_price),
                q.owner.get_full_name() or q.owner.username,
                q.created_at.strftime('%d %b %Y'),
            ]
            ws.append(row_data)
            for cell in ws[row_num]:
                cell.border = border
                cell.alignment = align_center if cell.column in [1, 3, 6, 7, 8, 9, 11] else align_left

        widths = {'A': 16, 'B': 28, 'C': 15, 'D': 15, 'E': 15, 'F': 16, 'G': 16, 'H': 16, 'I': 18, 'J': 20, 'K': 16}
        for col, w in widths.items():
            ws.column_dimensions[col].width = w
        ws.auto_filter.ref = f"A4:K{ws.max_row}"

        response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        response['Content-Disposition'] = f'attachment; filename="quotations_export_{tz.now().strftime("%Y%m%d")}.xlsx"'
        wb.save(response)
        return response

