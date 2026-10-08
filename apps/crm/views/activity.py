"""
Activity and Reminder views with RBAC enforcement.
Includes full CRUD: Create, Read (List), Update, Delete.
"""
from django.contrib import messages
from django.contrib.auth.mixins import LoginRequiredMixin
from django.views.generic import ListView, CreateView, UpdateView, DeleteView
from django.views import View
from django.urls import reverse_lazy
from django.shortcuts import get_object_or_404, redirect

from apps.core.mixins import OwnerRequiredMixin, OwnerOrAdminMixin, SetOwnerMixin, is_admin, AdminOnlyMixin
from apps.crm.models import Activity, Reminder
from apps.crm.forms import ActivityForm, ReminderForm


class ActivityListView(OwnerRequiredMixin, ListView):
    model = Activity
    template_name = 'crm/activities/activity_list.html'
    context_object_name = 'activities'
    paginate_by = 20

    def get_queryset(self):
        qs = super().get_queryset().select_related('lead__client', 'owner')
        search = self.request.GET.get('q', '').strip()
        activity_type = self.request.GET.get('type', '').strip()
        company = self.request.GET.get('company', '').strip()
        start_date = self.request.GET.get('start_date', '').strip()
        end_date = self.request.GET.get('end_date', '').strip()
        
        from apps.core.date_utils import parse_date_safe
        psd = parse_date_safe(start_date)
        ped = parse_date_safe(end_date)
        
        if search:
            from django.db.models import Q
            matched_qs = qs.filter(
                Q(lead__client__company_name__icontains=search) |
                Q(description__icontains=search) |
                Q(result__icontains=search)
            )
            filtered_qs = matched_qs
            if activity_type:
                filtered_qs = filtered_qs.filter(activity_type=activity_type)
            if company:
                filtered_qs = filtered_qs.filter(lead__client__company_name__icontains=company)
                
            date_filtered_qs = filtered_qs
            if psd:
                date_filtered_qs = date_filtered_qs.filter(created_at__date__gte=psd)
            if ped:
                date_filtered_qs = date_filtered_qs.filter(created_at__date__lte=ped)
                
            if date_filtered_qs.exists():
                return date_filtered_qs
            elif filtered_qs.exists():
                return filtered_qs
            return matched_qs
            
        if activity_type:
            qs = qs.filter(activity_type=activity_type)
        if company:
            qs = qs.filter(lead__client__company_name__icontains=company)
        if psd:
            qs = qs.filter(created_at__date__gte=psd)
        if ped:
            qs = qs.filter(created_at__date__lte=ped)
            
        return qs

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['search_query'] = self.request.GET.get('q', '')
        context['current_type'] = self.request.GET.get('type', '')
        context['start_date'] = self.request.GET.get('start_date', '')
        context['end_date'] = self.request.GET.get('end_date', '')
        context['company_query'] = self.request.GET.get('company', '')
        context['type_choices'] = Activity.ActivityType.choices
        return context


class ActivityCreateView(LoginRequiredMixin, SetOwnerMixin, CreateView):
    model = Activity
    form_class = ActivityForm
    template_name = 'crm/activities/activity_edit.html'

    def get_success_url(self):
        return self.object.lead.get_absolute_url()

    def get_initial(self):
        initial = super().get_initial()
        lead_id = self.request.GET.get('lead')
        if lead_id:
            initial['lead'] = lead_id
        return initial

    def get_form_kwargs(self):
        kwargs = super().get_form_kwargs()
        kwargs['user'] = self.request.user
        return kwargs

    def form_valid(self, form):
        messages.success(self.request, 'Aktivitas berhasil dicatat.')
        return super().form_valid(form)


class ActivityUpdateView(OwnerOrAdminMixin, UpdateView):
    model = Activity
    form_class = ActivityForm
    template_name = 'crm/activities/activity_edit.html'
    success_url = reverse_lazy('crm:activity-list')

    def get_form_kwargs(self):
        kwargs = super().get_form_kwargs()
        kwargs['user'] = self.request.user
        return kwargs

    def form_valid(self, form):
        messages.success(self.request, 'Aktivitas berhasil diperbarui.')
        return super().form_valid(form)


class ActivityDeleteView(AdminOnlyMixin, DeleteView):
    model = Activity
    template_name = 'crm/activities/activity_delete.html'
    success_url = reverse_lazy('crm:activity-list')

    def form_valid(self, form):
        messages.success(self.request, 'Aktivitas berhasil dihapus.')
        return super().form_valid(form)


class ReminderListView(OwnerRequiredMixin, ListView):
    model = Reminder
    template_name = 'crm/activities/reminder_list.html'
    context_object_name = 'reminders'
    paginate_by = 20

    def get_queryset(self):
        qs = super().get_queryset()
        show = self.request.GET.get('show', 'active')
        if show == 'completed':
            qs = qs.filter(completed=True)
        elif show == 'all':
            pass
        else:
            qs = qs.filter(completed=False)
        return qs

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['current_show'] = self.request.GET.get('show', 'active')
        return context


class ReminderCreateView(LoginRequiredMixin, SetOwnerMixin, CreateView):
    model = Reminder
    form_class = ReminderForm
    template_name = 'crm/activities/reminder_edit.html'
    success_url = reverse_lazy('crm:reminder-list')

    def form_valid(self, form):
        messages.success(self.request, 'Reminder berhasil dibuat.')
        return super().form_valid(form)


class ReminderUpdateView(OwnerOrAdminMixin, UpdateView):
    model = Reminder
    form_class = ReminderForm
    template_name = 'crm/activities/reminder_edit.html'
    success_url = reverse_lazy('crm:reminder-list')

    def form_valid(self, form):
        messages.success(self.request, 'Reminder berhasil diperbarui.')
        return super().form_valid(form)


class ReminderDeleteView(AdminOnlyMixin, DeleteView):
    model = Reminder
    template_name = 'crm/activities/reminder_delete.html'
    success_url = reverse_lazy('crm:reminder-list')

    def form_valid(self, form):
        messages.success(self.request, 'Reminder berhasil dihapus.')
        return super().form_valid(form)


class ReminderToggleCompleteView(LoginRequiredMixin, View):
    """Toggle the completed status of a reminder."""

    def post(self, request, pk):
        reminder = get_object_or_404(Reminder, pk=pk)
        if not is_admin(request.user) and reminder.owner != request.user:
            messages.error(request, 'Anda tidak memiliki akses.')
            return redirect('crm:reminder-list')

        reminder.completed = not reminder.completed
        reminder.save()
        status_text = 'selesai' if reminder.completed else 'aktif kembali'
        messages.success(request, f'Reminder "{reminder.title}" ditandai {status_text}.')
        return redirect('crm:reminder-list')

from django.http import JsonResponse
from django.utils import timezone

from datetime import datetime, timedelta

def due_reminders_api(request):
    if not request.user.is_authenticated:
        return JsonResponse({'reminders': []})
    
    now = timezone.localtime(timezone.now())
    today = now.date()
    
    reminders = Reminder.objects.filter(
        owner=request.user, 
        completed=False, 
        due_date__lte=today
    )
    
    due_list = []
    for r in reminders:
        if r.due_date < today:
            due_list.append({'id': r.id, 'title': r.title})
        elif r.due_date == today:
            if not r.reminder_time:
                due_list.append({'id': r.id, 'title': r.title})
            else:
                # Combine date and time, and make it aware in current timezone
                reminder_dt = timezone.make_aware(datetime.combine(r.due_date, r.reminder_time))
                notify_time = reminder_dt - timedelta(minutes=r.notify_before)
                if now >= notify_time:
                    due_list.append({'id': r.id, 'title': r.title})
                
    return JsonResponse({'reminders': due_list})


class ActivityExportExcelView(OwnerRequiredMixin, ListView):
    """Export Activities to Excel (.xlsx) with styling."""
    model = Activity

    def get(self, request, *args, **kwargs):
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from django.http import HttpResponse
        from django.utils import timezone as tz

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Activities Export"

        title_font = Font(size=16, bold=True, color="000000")
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="343A40", end_color="343A40", fill_type="solid")
        border = Border(
            left=Side(style='thin'), right=Side(style='thin'),
            top=Side(style='thin'), bottom=Side(style='thin')
        )
        align_center = Alignment(horizontal="center", vertical="center")
        align_left = Alignment(horizontal="left", vertical="center")

        ws.merge_cells('A1:H1')
        ws['A1'].value = "DATA AKTIVITAS SALES"
        ws['A1'].font = title_font
        ws['A1'].alignment = align_center

        ws.merge_cells('A2:H2')
        ws['A2'].value = f"Diekspor pada: {tz.now().strftime('%d %b %Y')}"
        ws['A2'].alignment = align_center
        ws.append([])

        headers = [
            'Perusahaan', 'Jenis', 'Deskripsi', 'Hasil',
            'Follow Up Selanjutnya', 'PIC', 'Tanggal', 'Lead'
        ]
        ws.append(headers)
        for cell in ws[4]:
            cell.font = header_font
            cell.fill = header_fill
            cell.border = border
            cell.alignment = align_center

        qs = self.get_queryset().select_related('lead__client', 'owner')
        search = request.GET.get('q', '')
        activity_type = request.GET.get('type', '')
        company = request.GET.get('company', '')
        start_date = request.GET.get('start_date', '')
        end_date = request.GET.get('end_date', '')

        if search:
            from django.db.models import Q
            qs = qs.filter(
                Q(lead__client__company_name__icontains=search) |
                Q(description__icontains=search) |
                Q(result__icontains=search)
            )
        if activity_type:
            qs = qs.filter(activity_type=activity_type)
        if company:
            qs = qs.filter(lead__client__company_name__icontains=company)
        if start_date:
            qs = qs.filter(created_at__gte=f"{start_date} 00:00:00")
        if end_date:
            qs = qs.filter(created_at__lte=f"{end_date} 23:59:59")

        for row_num, act in enumerate(qs, start=5):
            row_data = [
                act.lead.client.company_name,
                act.get_activity_type_display(),
                act.description[:100],
                act.result[:100] if act.result else '-',
                act.next_followup.strftime('%d %b %Y') if act.next_followup else '-',
                act.owner.get_full_name() or act.owner.username,
                act.created_at.strftime('%d %b %Y'),
                act.lead.lead_id,
            ]
            ws.append(row_data)
            for cell in ws[row_num]:
                cell.border = border
                cell.alignment = align_center if cell.column in [2, 5, 6, 7, 8] else align_left

        widths = {'A': 25, 'B': 15, 'C': 40, 'D': 40, 'E': 20, 'F': 20, 'G': 18, 'H': 12}
        for col, w in widths.items():
            ws.column_dimensions[col].width = w
        ws.auto_filter.ref = f"A4:H{ws.max_row}"

        response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        response['Content-Disposition'] = f'attachment; filename="activities_export_{tz.now().strftime("%Y%m%d")}.xlsx"'
        wb.save(response)
        return response


class ReminderExportExcelView(OwnerRequiredMixin, ListView):
    """Export Reminders to Excel (.xlsx) with styling."""
    model = Reminder

    def get(self, request, *args, **kwargs):
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from django.http import HttpResponse
        from django.utils import timezone as tz

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Reminders Export"

        title_font = Font(size=16, bold=True, color="000000")
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="343A40", end_color="343A40", fill_type="solid")
        border = Border(
            left=Side(style='thin'), right=Side(style='thin'),
            top=Side(style='thin'), bottom=Side(style='thin')
        )
        align_center = Alignment(horizontal="center", vertical="center")
        align_left = Alignment(horizontal="left", vertical="center")

        ws.merge_cells('A1:F1')
        ws['A1'].value = "DATA REMINDER"
        ws['A1'].font = title_font
        ws['A1'].alignment = align_center

        ws.merge_cells('A2:F2')
        ws['A2'].value = f"Diekspor pada: {tz.now().strftime('%d %b %Y')}"
        ws['A2'].alignment = align_center
        ws.append([])

        headers = ['Judul', 'Deskripsi', 'Tanggal', 'Waktu', 'Status', 'PIC']
        ws.append(headers)
        for cell in ws[4]:
            cell.font = header_font
            cell.fill = header_fill
            cell.border = border
            cell.alignment = align_center

        qs = self.get_queryset()
        show = request.GET.get('show', 'active')
        if show == 'completed':
            qs = qs.filter(completed=True)
        elif show == 'all':
            pass
        else:
            qs = qs.filter(completed=False)

        for row_num, rem in enumerate(qs, start=5):
            row_data = [
                rem.title,
                rem.description[:100] if rem.description else '-',
                rem.due_date.strftime('%d %b %Y'),
                rem.reminder_time.strftime('%H:%M') if rem.reminder_time else '-',
                'Selesai' if rem.completed else 'Aktif',
                rem.owner.get_full_name() or rem.owner.username,
            ]
            ws.append(row_data)
            for cell in ws[row_num]:
                cell.border = border
                cell.alignment = align_center if cell.column in [3, 4, 5, 6] else align_left

        widths = {'A': 30, 'B': 40, 'C': 18, 'D': 12, 'E': 12, 'F': 20}
        for col, w in widths.items():
            ws.column_dimensions[col].width = w
        ws.auto_filter.ref = f"A4:F{ws.max_row}"

        response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        response['Content-Disposition'] = f'attachment; filename="reminders_export_{tz.now().strftime("%Y%m%d")}.xlsx"'
        wb.save(response)
        return response

