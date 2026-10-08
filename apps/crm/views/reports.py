import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from django.shortcuts import render
from django.http import HttpResponse
from django.contrib.auth.decorators import login_required
from django.utils import timezone
from django.db.models import Sum, Count, Avg, Q
from django.core.paginator import Paginator
from django.contrib.auth import get_user_model

from apps.crm.models import Lead, Client, Contract, Quotation, Activity

User = get_user_model()

@login_required
def sales_report_view(request):
    start_date = request.GET.get('start_date', '')
    end_date = request.GET.get('end_date', '')
    owner_id = request.GET.get('owner', '')
    status_filter = request.GET.get('status', '')
    source_filter = request.GET.get('source', '')
    industry_filter = request.GET.get('industry', '')
    search_query = request.GET.get('q', '').strip()

    today = timezone.localdate()
    default_start = today - datetime.timedelta(days=30)

    queryset = Lead.objects.select_related('owner', 'client').order_by('-created_at')

    # Date filter
    if start_date:
        queryset = queryset.filter(created_at__date__gte=start_date)
    else:
        queryset = queryset.filter(created_at__date__gte=default_start)
        start_date = default_start.isoformat()

    if end_date:
        queryset = queryset.filter(created_at__date__lte=end_date)
    else:
        end_date = today.isoformat()

    if owner_id:
        queryset = queryset.filter(owner_id=owner_id)
    if status_filter:
        queryset = queryset.filter(status=status_filter)
    if source_filter:
        queryset = queryset.filter(lead_source=source_filter)
    if industry_filter:
        queryset = queryset.filter(client__industry=industry_filter)
    if search_query:
        queryset = queryset.filter(
            Q(client__company_name__icontains=search_query) |
            Q(client__contact_person__icontains=search_query) |
            Q(owner__username__icontains=search_query) |
            Q(owner__first_name__icontains=search_query) |
            Q(owner__last_name__icontains=search_query)
        )

    all_leads = list(queryset)
    total_leads_count = len(all_leads)
    total_pipeline_value = sum([float(l.estimated_value or 0) for l in all_leads])

    deals_won_count = 0
    deals_won_value = 0.0
    deals_lost_count = 0
    deals_lost_value = 0.0
    total_closed_days = 0.0
    closed_leads_count = 0

    funnel_counts = {code: 0 for code, label in Lead.Status.choices}
    funnel_values = {code: 0.0 for code, label in Lead.Status.choices}
    source_distribution = {}
    pic_performance = {}

    for lead in all_leads:
        val = float(lead.estimated_value or 0)
        st = lead.status
        
        # Funnel stage
        if st in funnel_counts:
            funnel_counts[st] += 1
            funnel_values[st] += val

        # Won & Lost Counters
        if st == 'WON':
            deals_won_count += 1
            deals_won_value += val
        elif st == 'LOST':
            deals_lost_count += 1
            deals_lost_value += val

        # Closing duration
        if lead.closed_at and lead.created_at:
            dur_days = (lead.closed_at - lead.created_at).total_seconds() / 86400
            total_closed_days += max(0, dur_days)
            closed_leads_count += 1

        # Source breakdown
        src_label = lead.get_lead_source_display()
        source_distribution[src_label] = source_distribution.get(src_label, 0) + 1

        # PIC Sales Breakdown
        pic_name = (lead.owner.get_full_name() or lead.owner.username) if lead.owner else 'Tanpa PIC'
        if pic_name not in pic_performance:
            pic_performance[pic_name] = {
                'name': pic_name,
                'total_leads': 0,
                'total_val': 0.0,
                'won_count': 0,
                'won_val': 0.0,
                'lost_count': 0
            }
        pic_performance[pic_name]['total_leads'] += 1
        pic_performance[pic_name]['total_val'] += val
        if st == 'WON':
            pic_performance[pic_name]['won_count'] += 1
            pic_performance[pic_name]['won_val'] += val
        elif st == 'LOST':
            pic_performance[pic_name]['lost_count'] += 1

    # Calculate Win Rate & Avg Cycle
    closed_total = deals_won_count + deals_lost_count
    win_rate = round((deals_won_count / closed_total) * 100, 1) if closed_total > 0 else (round((deals_won_count / total_leads_count) * 100, 1) if total_leads_count > 0 else 0.0)
    avg_closing_days = round(total_closed_days / closed_leads_count, 1) if closed_leads_count > 0 else 0.0

    # Formatted strings
    total_pipeline_value_display = f"Rp {int(total_pipeline_value):,}".replace(',', '.')
    deals_won_value_display = f"Rp {int(deals_won_value):,}".replace(',', '.')
    deals_lost_value_display = f"Rp {int(deals_lost_value):,}".replace(',', '.')

    for lead in all_leads:
        lead.estimated_value_formatted = f"Rp {int(lead.estimated_value or 0):,}".replace(',', '.')
        lead.quotations_count = lead.quotations.count() if hasattr(lead, 'quotations') else 0

    # PIC Leaderboard list sorted by Won Value desc
    pic_leaderboard = []
    for p_name, p_data in pic_performance.items():
        p_data['win_rate'] = round((p_data['won_count'] / p_data['total_leads']) * 100, 1) if p_data['total_leads'] > 0 else 0.0
        p_data['won_val_formatted'] = f"Rp {int(p_data['won_val']):,}".replace(',', '.')
        p_data['total_val_formatted'] = f"Rp {int(p_data['total_val']):,}".replace(',', '.')
        pic_leaderboard.append(p_data)
    pic_leaderboard.sort(key=lambda x: x['won_val'], reverse=True)

    # Funnel Chart Data
    funnel_labels = [label for code, label in Lead.Status.choices]
    funnel_data = [funnel_counts[code] for code, label in Lead.Status.choices]

    # Quotations summary for the period
    quotations_qs = Quotation.objects.filter(created_at__date__gte=start_date, created_at__date__lte=end_date)
    quotations_count = quotations_qs.count()
    quotations_total_value = sum([float(q.manual_total_price or 0) for q in quotations_qs])
    quotations_total_value_display = f"Rp {int(quotations_total_value):,}".replace(',', '.')

    # Pagination
    paginator = Paginator(all_leads, 25)
    page_number = request.GET.get('page')
    page_obj = paginator.get_page(page_number)

    # Sales PIC Users list for filter dropdown
    sales_users = User.objects.filter(is_active=True).order_by('first_name', 'username')

    context = {
        'page_obj': page_obj,
        'total_leads_count': total_leads_count,
        'total_pipeline_value': total_pipeline_value,
        'total_pipeline_value_display': total_pipeline_value_display,
        'deals_won_count': deals_won_count,
        'deals_won_value': deals_won_value,
        'deals_won_value_display': deals_won_value_display,
        'deals_lost_count': deals_lost_count,
        'deals_lost_value': deals_lost_value,
        'deals_lost_value_display': deals_lost_value_display,
        'win_rate': win_rate,
        'avg_closing_days': avg_closing_days,
        'quotations_count': quotations_count,
        'quotations_total_value': quotations_total_value,
        'quotations_total_value_display': quotations_total_value_display,
        'funnel_labels': funnel_labels,
        'funnel_data': funnel_data,
        'source_labels': list(source_distribution.keys()),
        'source_values': list(source_distribution.values()),
        'pic_leaderboard': pic_leaderboard,
        'sales_users': sales_users,
        'status_choices': Lead.Status.choices,
        'source_choices': Lead.LeadSource.choices,
        'industry_choices': Client.IndustryChoices.choices,
        'start_date': start_date,
        'end_date': end_date,
        'owner_id': owner_id,
        'status_filter': status_filter,
        'source_filter': source_filter,
        'industry_filter': industry_filter,
        'search_query': search_query,
    }
    return render(request, 'crm/reports/sales_report.html', context)


@login_required
def sales_report_export_excel(request):
    start_date = request.GET.get('start_date', '')
    end_date = request.GET.get('end_date', '')
    owner_id = request.GET.get('owner', '')
    status_filter = request.GET.get('status', '')
    source_filter = request.GET.get('source', '')
    industry_filter = request.GET.get('industry', '')
    search_query = request.GET.get('q', '').strip()

    queryset = Lead.objects.select_related('owner', 'client').order_by('-created_at')

    if start_date:
        queryset = queryset.filter(created_at__date__gte=start_date)
    if end_date:
        queryset = queryset.filter(created_at__date__lte=end_date)
    if owner_id:
        queryset = queryset.filter(owner_id=owner_id)
    if status_filter:
        queryset = queryset.filter(status=status_filter)
    if source_filter:
        queryset = queryset.filter(lead_source=source_filter)
    if industry_filter:
        queryset = queryset.filter(client__industry=industry_filter)
    if search_query:
        queryset = queryset.filter(
            Q(client__company_name__icontains=search_query) |
            Q(client__contact_person__icontains=search_query) |
            Q(owner__username__icontains=search_query) |
            Q(owner__first_name__icontains=search_query) |
            Q(owner__last_name__icontains=search_query)
        )

    all_leads = list(queryset)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Laporan Performa Sales"

    title_font = Font(name='Calibri', size=16, bold=True, color='1F2937')
    subtitle_font = Font(name='Calibri', size=10, italic=True, color='6B7280')
    header_font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='DC2626', end_color='DC2626', fill_type='solid')
    kpi_val_font = Font(name='Calibri', size=13, bold=True, color='1F2937')
    kpi_fill = PatternFill(start_color='F3F4F6', end_color='F3F4F6', fill_type='solid')
    data_font = Font(name='Calibri', size=10)

    thin_border = Border(
        left=Side(style='thin', color='E5E7EB'),
        right=Side(style='thin', color='E5E7EB'),
        top=Side(style='thin', color='E5E7EB'),
        bottom=Side(style='thin', color='E5E7EB')
    )

    # Title Block
    ws['A1'] = "LAPORAN PIPELINE & PERFORMA SALES CRM"
    ws['A1'].font = title_font
    ws['A2'] = f"Periode: {start_date or 'Semua'} s/d {end_date or 'Semua'} | Diekspor: {timezone.localtime().strftime('%d/%m/%Y %H:%M:%S')} oleh {request.user.get_full_name() or request.user.username}"
    ws['A2'].font = subtitle_font

    total_leads = len(all_leads)
    total_val = sum([float(l.estimated_value or 0) for l in all_leads])
    won_leads = [l for l in all_leads if l.status == 'WON']
    won_val = sum([float(l.estimated_value or 0) for l in won_leads])
    win_rate = f"{round((len(won_leads) / total_leads) * 100, 1)}%" if total_leads > 0 else "0%"

    # KPI Banner
    kpis = [
        ("TOTAL LEADS", f"{total_leads} Prospek", "A4", "B5"),
        ("TOTAL PIPELINE", f"Rp {total_val:,.0f}", "C4", "D5"),
        ("DEALS WON", f"{len(won_leads)} Leads", "E4", "F5"),
        ("NILAI CLOSING WON", f"Rp {won_val:,.0f}", "G4", "H5"),
        ("WIN RATE", win_rate, "I4", "J5"),
    ]
    for title, val, c1, c2 in kpis:
        col_start, row_start = c1[0], int(c1[1])
        col_end, row_end = c2[0], int(c2[1])
        ws.merge_cells(f"{c1}:{c2}")
        top_left = ws[f"{col_start}{row_start}"]
        top_left.value = f"{title}\n{val}"
        top_left.font = kpi_val_font
        top_left.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        top_left.fill = kpi_fill
        for r in range(row_start, row_end + 1):
            for col_idx in [openpyxl.utils.column_index_from_string(col_start), openpyxl.utils.column_index_from_string(col_end)]:
                ws.cell(row=r, column=col_idx).border = thin_border

    # Headers
    headers = [
        "No", "ID Lead", "Nama Perusahaan / Klien", "Kontak Person (PIC)", "No. Telepon",
        "PIC Sales", "Sumber Lead", "Industri", "Estimasi Nilai (Rp)", "Status / Tahapan",
        "Tanggal Dibuat", "Tanggal Closing", "Respon Cepat", "Decision Maker", "Ada Budget"
    ]
    header_row = 7
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=header_row, column=col_idx)
        cell.value = header
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = thin_border

    ws.row_dimensions[header_row].height = 28

    # Data Rows
    start_row = 8
    for idx, lead in enumerate(all_leads, 1):
        est_val = float(lead.estimated_value or 0)
        row_values = [
            idx,
            lead.lead_id,
            lead.client.company_name if lead.client else '',
            lead.client.contact_person if lead.client else '',
            lead.client.phone if lead.client else '',
            (lead.owner.get_full_name() or lead.owner.username) if lead.owner else '',
            lead.get_lead_source_display(),
            lead.client.get_industry_display() if (lead.client and lead.client.industry) else '',
            est_val,
            lead.get_status_display(),
            lead.created_at.strftime('%d/%m/%Y %H:%M') if lead.created_at else '',
            lead.closed_at.strftime('%d/%m/%Y %H:%M') if lead.closed_at else '',
            lead.get_responds_fast_display(),
            lead.get_decision_maker_display(),
            lead.get_budget_available_display(),
        ]
        curr_row = start_row + idx - 1
        for col_idx, val in enumerate(row_values, 1):
            cell = ws.cell(row=curr_row, column=col_idx)
            cell.value = val
            cell.font = data_font
            cell.border = thin_border

            if col_idx in [1, 2, 7, 10, 11, 12, 13, 14, 15]:
                cell.alignment = Alignment(horizontal='center', vertical='center')
            elif col_idx == 9:
                cell.alignment = Alignment(horizontal='right', vertical='center')
                cell.number_format = '#,##0'
            else:
                cell.alignment = Alignment(horizontal='left', vertical='center')

            # Highlight status
            if col_idx == 10:
                if val == 'Menang / Berhasil':
                    cell.fill = PatternFill(start_color='DCFCE7', end_color='DCFCE7', fill_type='solid')
                    cell.font = Font(name='Calibri', size=10, bold=True, color='166534')
                elif val == 'Kalah / Gagal':
                    cell.fill = PatternFill(start_color='FEE2E2', end_color='FEE2E2', fill_type='solid')
                    cell.font = Font(name='Calibri', size=10, bold=True, color='991B1B')

    # Auto-adjust column widths based only on header and data
    for col_idx in range(1, len(headers) + 1):
        col_letter = openpyxl.utils.get_column_letter(col_idx)
        if col_idx == 1:
            ws.column_dimensions[col_letter].width = 6
        else:
            max_len = max(len(str(ws.cell(row=r, column=col_idx).value or '')) for r in range(header_row, ws.max_row + 1)) if ws.max_row >= header_row else len(headers[col_idx-1])
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    filename = f"Laporan_Sales_{timezone.localdate().strftime('%Y%m%d')}.xlsx"
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    wb.save(response)
    return response
