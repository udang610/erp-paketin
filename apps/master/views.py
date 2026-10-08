from django.shortcuts import render, redirect
from django.contrib.auth.decorators import login_required
from django.contrib import messages

# Map model names to Forms and Models
def get_model_and_form(model_name):
    from .models import Bank, Coverage, Service, Price, Vehicle, Customer
    from .forms import BankForm, CoverageForm, ServiceForm, PriceForm, VehicleForm, BranchForm, ClientForm, VendorForm, CustomerForm
    
    from apps.organizations.models import Branch
    from apps.crm.models import Client
    from apps.finance.models import Vendor
    
    mapping = {
        'bank': (Bank, BankForm, 'master:banks'),
        'coverage': (Coverage, CoverageForm, 'master:coverage'),
        'service': (Service, ServiceForm, 'master:services'),
        'price': (Price, PriceForm, 'master:price'),
        'vehicle': (Vehicle, VehicleForm, 'master:vehicle'),
        'branch': (Branch, BranchForm, 'master:branches'),
        'customer': (Customer, CustomerForm, 'master:customer'),
        'vendor': (Vendor, VendorForm, 'master:vendors'),
    }
    return mapping.get(model_name, (None, None, None))

def _build_master_filters(request, model_name):
    """Return (filter_fields for template, ORM filter kwargs) per master model."""
    from apps.organizations.models import Branch
    from .models import Vehicle, Customer, Service, Coverage

    def branches():
        return [(str(b.pk), f"{b.name} ({b.code})") for b in Branch.objects.order_by('name')]

    active_opts = [('1', 'Aktif'), ('0', 'Nonaktif')]
    yes_no = [('1', 'Ya'), ('0', 'Tidak')]
    # (param, label, type, options_callable, lookup, cast)
    cfg = {
        'bank': [
            ('branch', 'Cabang Terkait', 'select', branches, 'branch_id', None),
            ('active', 'Status', 'select', lambda: active_opts, 'is_active', 'bool'),
        ],
        'branch': [
            ('warehouse', 'Memiliki Gudang', 'select', lambda: yes_no, 'has_warehouse', 'bool'),
            ('timezone', 'Zona Waktu', 'select', lambda: [(t, t) for t in Branch.objects.values_list('timezone', flat=True).distinct().order_by('timezone')], 'timezone', None),
        ],
        'coverage': [
            ('province', 'Provinsi', 'select', lambda: [(p, p) for p in Coverage.objects.values_list('province', flat=True).distinct().order_by('province')], 'province', None),
            ('city', 'Kota / Kabupaten', 'text', None, 'city__icontains', None),
            ('covered', 'Status Jangkauan', 'select', lambda: [('1', 'Terjangkau'), ('0', 'Tidak Terjangkau')], 'is_covered', 'bool'),
            ('active', 'Status', 'select', lambda: active_opts, 'is_active', 'bool'),
        ],
        'customer': [
            ('branch', 'Cabang', 'select', branches, 'branch_id', None),
            ('industry', 'Tipe Industri', 'select', lambda: Customer.INDUSTRY_CHOICES, 'type_industry', None),
            ('account_type', 'Tipe Akun', 'select', lambda: Customer.ACCOUNT_TYPE_CHOICES, 'account_type', None),
            ('status', 'Status Customer', 'select', lambda: Customer.STATUS_CUSTOMER_CHOICES, 'status', None),
            ('payment_type', 'Tipe Pembayaran', 'select', lambda: Customer.PAYMENT_TYPE_CHOICES, 'payment_type', None),
            ('top', 'Term of Payment', 'select', lambda: Customer.TERM_CHOICES, 'term_of_payment', None),
        ],
        'vehicle': [
            ('branch', 'Cabang', 'select', branches, 'branch_id', None),
            ('vehicle_type', 'Tipe Kendaraan', 'select', lambda: [(t, t) for t in Vehicle.objects.values_list('vehicle_type', flat=True).distinct().order_by('vehicle_type')], 'vehicle_type', None),
            ('status', 'Status Armada', 'select', lambda: Vehicle.STATUS_CHOICES, 'status', None),
        ],
        'price': [
            ('origin_city', 'Kota Asal', 'text', None, 'origin__city__icontains', None),
            ('destination_city', 'Kota Tujuan', 'text', None, 'destination__city__icontains', None),
            ('service', 'Layanan', 'select', lambda: [(str(s.pk), f"{s.name} ({s.code})") for s in Service.objects.order_by('name')], 'service_id', None),
            ('active', 'Status', 'select', lambda: active_opts, 'is_active', 'bool'),
        ],
        'service': [
            ('active', 'Status', 'select', lambda: active_opts, 'is_active', 'bool'),
        ],
        'vendor': [
            ('transport', 'Moda Transportasi Utama', 'select', lambda: [('DARAT', 'Darat'), ('LAUT', 'Laut'), ('UDARA', 'Udara'), ('KERETA', 'Kereta')], 'primary_transport_mode', None),
            ('delivery_type', 'Tipe Pengiriman', 'select', lambda: [('Door to Door', 'Door to Door'), ('Port to Port', 'Port to Port'), ('Door to Port', 'Door to Port'), ('Port to Door', 'Port to Door')], 'delivery_type', None),
            ('branch', 'Cabang Terkait', 'select', branches, 'branch_id', None),
            ('active', 'Status', 'select', lambda: active_opts, 'is_active', 'bool'),
        ],
    }
    fields, q = [], {}
    for param, label, ftype, opts, lookup, cast in cfg.get(model_name, []):
        value = request.GET.get(param, '').strip()
        fields.append({'name': param, 'label': label, 'type': ftype, 'options': opts() if opts else [], 'value': value})
        if value:
            q[lookup] = (value == '1') if cast == 'bool' else value
    
    from apps.core.date_utils import parse_date_safe
    start_date = request.GET.get('start_date', '').strip()
    end_date = request.GET.get('end_date', '').strip()
    
    # Master reference tables (bank, branch, service, price, coverage, vehicle, vendor) are static reference data
    # Only show date range inputs for transactional master tables like customer
    if model_name in ['customer']:
        date_fields = [
            {'name': 'start_date', 'label': 'Tanggal Mulai', 'type': 'date', 'options': [], 'value': start_date},
            {'name': 'end_date', 'label': 'Tanggal Akhir', 'type': 'date', 'options': [], 'value': end_date}
        ]
    else:
        date_fields = []

    psd = parse_date_safe(start_date)
    ped = parse_date_safe(end_date)
    if psd:
        q['created_at__date__gte'] = psd
    if ped:
        q['created_at__date__lte'] = ped

    return fields, date_fields, q

MASTER_MENU_FOLDERS = {
    'bank': ('banks', 'bank'),
    'branch': ('branches', 'branch'),
    'coverage': ('coverage', 'coverage'),
    'customer': ('customers', 'customer'),
    'vehicle': ('vehicles', 'vehicle'),
    'price': ('prices', 'price'),
    'service': ('services', 'service'),
    'vendor': ('vendors', 'vendor'),
}

def get_master_template(model_name, action):
    folder, prefix = MASTER_MENU_FOLDERS.get(model_name, ('', model_name))
    if folder:
        return f'master/{folder}/{prefix}_{action}.html'
    return f'master/master_{action}.html'

@login_required
def master_list(request, model_name, title, icon):
    Model, _, _ = get_model_and_form(model_name)
    
    search_query = request.GET.get('q', '').strip()
    items = []
    filter_fields = []
    active_filters = False
    if Model:
        from django.db.models import Q
        items = Model.objects.all()

        # Universal search: one keyword matched against every key text column
        search_fields_map = {
            'coverage': ['city', 'district', 'province', 'tlc', 'postal_code'],
            'service': ['code', 'name', 'description'],
            'customer': ['customer_code', 'name', 'city', 'district', 'pic_account', 'pic_phone', 'email', 'npwp', 'branch__name', 'branch__code', 'sales__username', 'sales__first_name'],
            'vendor': ['name', 'contact_person', 'phone', 'email', 'city', 'primary_transport_mode', 'service_type', 'npwp'],
            'branch': ['code', 'name', 'address'],
            'vehicle': ['plate_number', 'vehicle_type', 'brand_model', 'branch__name', 'branch__code'],
            'price': ['origin__city', 'origin__district', 'origin__tlc', 'destination__city', 'destination__district', 'destination__tlc', 'service__name', 'service__code'],
            'bank': ['code', 'name', 'account_number', 'account_name', 'address', 'branch__name', 'branch__code'],
        }
        # Model-specific popup filters
        filter_fields, date_fields, filter_q = _build_master_filters(request, model_name)

        if search_query:
            q_obj = Q()
            for f in search_fields_map.get(model_name, []):
                q_obj |= Q(**{f'{f}__icontains': search_query})
            if model_name == 'price':
                items = items.select_related('origin', 'destination', 'service')
            
            # Universal keyword matching
            matched_items = items.filter(q_obj)
            
            if filter_q:
                refined_items = matched_items.filter(**filter_q)
                if refined_items.exists():
                    items = refined_items
                    active_filters = True
                else:
                    # If date filter caused 0 results on a keyword search, ignore the date filter
                    non_date_q = {k: v for k, v in filter_q.items() if not k.startswith('created_at')}
                    if non_date_q:
                        items = matched_items.filter(**non_date_q)
                        active_filters = True
                    else:
                        items = matched_items
                        active_filters = False
            else:
                items = matched_items
        else:
            if filter_q:
                items_filtered = items.filter(**filter_q)
                if items_filtered.exists():
                    items = items_filtered
                    active_filters = True
                else:
                    # If date filter caused 0 results on non-search filter, fallback to non-date filter!
                    non_date_q = {k: v for k, v in filter_q.items() if not k.startswith('created_at')}
                    if non_date_q:
                        items_fallback = items.filter(**non_date_q)
                        if items_fallback.exists():
                            items = items_fallback
                            active_filters = True
                        else:
                            items = items_filtered
                            active_filters = True
                    else:
                        items = items_filtered
                        active_filters = True
        
        if model_name == 'price':
            items = items.order_by('id')
        elif model_name == 'customer':
            items = items.order_by('-id')
        elif model_name == 'coverage':
            items = items.order_by('id')
        else:
            items = items.order_by('-id')

        from django.core.paginator import Paginator
        from django.core.cache import cache
        paginator = Paginator(items, 50)
        
        if not search_query and not active_filters:
            cache_key = f'master_total_count_{model_name}'
            cached_count = cache.get(cache_key)
            if cached_count is None:
                cached_count = paginator.count
                cache.set(cache_key, cached_count, 300)
            paginator._count = cached_count
            total_count = cached_count
        else:
            total_count = paginator.count

        page_number = request.GET.get('page', 1)
        page_obj = paginator.get_page(page_number)

        if model_name == 'price':
            current_items = list(page_obj.object_list.select_related('origin', 'destination', 'service', 'created_by'))
        elif model_name == 'customer':
            current_items = list(page_obj.object_list.select_related('branch', 'sales', 'created_by'))
        elif model_name == 'coverage':
            current_items = list(page_obj.object_list.select_related('created_by'))
        else:
            current_items = list(page_obj.object_list.select_related('created_by') if hasattr(Model, 'created_by') else page_obj.object_list)
    else:
        page_obj = None
        total_count = 0
        current_items = []
    
    is_implemented = Model is not None
    
    header_mapping = {
        'coverage': ['Provinsi', 'Kota/Kabupaten', 'Kecamatan', 'TLC', 'Status', 'Dibuat Oleh', 'Waktu Dibuat'],
        'service': ['Kode', 'Nama Layanan', 'Pembagi (Divisor)', 'Status', 'Dibuat Oleh', 'Waktu Dibuat'],
        'customer': ['Kode Kustomer', 'Nama Kustomer', 'Cabang', 'Nama PIC', 'Telepon', 'Kota', 'Sales / CRM', 'Dibuat Oleh', 'Waktu Dibuat'],
        'vendor': ['Nama Vendor', 'Nama PIC', 'Moda Utama', 'Telepon', 'Kota', 'Dibuat Oleh', 'Waktu Dibuat'],
        'branch': ['Kode Cabang', 'Nama Cabang', 'Alamat', 'Zona Waktu', 'Gudang', 'Dibuat Oleh', 'Waktu Dibuat'],
        'vehicle': ['No. Polisi', 'Tipe Kendaraan', 'Merk / Model', 'Cabang', 'Status', 'Dibuat Oleh', 'Waktu Dibuat'],
        'bank': ['Kode', 'Nama Bank', 'No. Rekening', 'Atas Nama', 'Cabang', 'Status', 'Dibuat Oleh', 'Waktu Dibuat'],
        'price': ['Asal (Origin)', 'Tujuan (Destination)', 'Layanan', 'Min. Berat', 'Harga / Kg', 'Estimasi (SLA)', 'Status', 'Dibuat Oleh', 'Waktu Dibuat'],
    }
    headers = header_mapping.get(model_name, ['Data (Nama / Deskripsi)', 'Dibuat Oleh', 'Waktu Dibuat'])
    table_data = []
    
    if Model:
        from apps.master.display import clean_branch_label
        for item in current_items:
            fields = []
            if model_name == 'coverage':
                import re
                clean_city = re.sub(r'(?i)^kota\s+', '', item.city)
                clean_city = re.sub(r'(?i)^kabupaten\s+', 'Kab. ', clean_city)
                fields = [item.province, clean_city, item.district, item.tlc or '-', 'Aktif' if item.is_active else 'Nonaktif']
            elif model_name == 'service':
                fields = [item.code, item.name, item.divisor, 'Aktif' if item.is_active else 'Nonaktif']
            elif model_name == 'customer':
                import re
                clean_city = re.sub(r'(?i)^kota\s+', '', item.city) if item.city else '-'
                clean_city = re.sub(r'(?i)^kabupaten\s+', 'Kab. ', clean_city) if clean_city != '-' else '-'
                sales_info = item.sales.get_full_name() or item.sales.username if item.sales else (item.crm_client.owner.username if hasattr(item, 'crm_client') and item.crm_client and item.crm_client.owner else '-')
                fields = [item.customer_code, item.name, clean_branch_label(item.branch), item.pic_account or '-', item.pic_phone or '-', clean_city, sales_info]
            elif model_name == 'vendor':
                import re
                clean_city = re.sub(r'(?i)^kota\s+', '', item.city) if item.city else '-'
                clean_city = re.sub(r'(?i)^kabupaten\s+', 'Kab. ', clean_city) if clean_city != '-' else '-'
                fields = [item.name, item.contact_person or '-', item.primary_transport_mode or '-', item.phone or '-', clean_city]
            elif model_name == 'branch':
                fields = [item.code, clean_branch_label(item), item.address or '-', item.timezone, 'Ya' if item.has_warehouse else 'Tidak']
            elif model_name == 'vehicle':
                fields = [item.plate_number, item.vehicle_type, item.brand_model, clean_branch_label(item.branch), 'Aktif' if getattr(item, 'is_active', True) else 'Nonaktif']
            elif model_name == 'bank':
                fields = [item.code, item.name, item.account_number or '-', item.account_name or '-', clean_branch_label(item.branch), 'Aktif' if getattr(item, 'is_active', True) else 'Nonaktif']
            elif model_name == 'price':
                import re
                orig_city = re.sub(r'(?i)^(kota|kabupaten|kab\.)\s+', '', item.origin.city).strip().title() if item.origin else '-'
                orig_tlc = f" ({item.origin.tlc})" if item.origin and item.origin.tlc else ''
                origin_label = f"{item.origin.district}, {orig_city}{orig_tlc}" if item.origin else '-'

                dest_city = re.sub(r'(?i)^(kota|kabupaten|kab\.)\s+', '', item.destination.city).strip().title() if item.destination else '-'
                dest_tlc = f" ({item.destination.tlc})" if item.destination and item.destination.tlc else ''
                dest_label = f"{item.destination.district}, {dest_city}{dest_tlc}" if item.destination else '-'

                service_label = f"{item.service.name} ({item.service.code})" if item.service else '-'
                min_w = f"{float(item.min_weight or 1):,.1f} Kg"
                p_kg = f"Rp {int(item.price_per_kg):,}".replace(',', '.') if item.price_per_kg is not None else '-'
                fields = [origin_label, dest_label, service_label, min_w, p_kg, item.estimated_days or '-', 'Aktif' if getattr(item, 'is_active', True) else 'Nonaktif']
            else:
                fields = [str(item)]
                
            # Tambahkan Dibuat Oleh & Waktu Dibuat
            creator_label = item.created_by.get_full_name() or item.created_by.username if getattr(item, 'created_by', None) else 'System'
            created_at_label = item.created_at.strftime('%d/%m/%Y %H:%M') if getattr(item, 'created_at', None) else '-'
            fields.append(creator_label)
            fields.append(created_at_label)

            table_data.append({
                'id': item.pk,
                'fields': fields
            })
            
    return render(request, get_master_template(model_name, 'list'), {
        'model_name': model_name,
        'title': title,
        'icon': icon,
        'headers': headers,
        'table_data': table_data,
        'is_implemented': is_implemented,
        'search_query': search_query,
        'filter_fields': filter_fields,
        'date_fields': date_fields,
        'active_filters': active_filters,
        'page_obj': page_obj,
        'total_count': total_count,
    })

@login_required
def master_create(request, model_name):
    Model, FormClass, redirect_url = get_model_and_form(model_name)
    if not FormClass:
        messages.error(request, f"Form untuk {model_name} belum tersedia.")
        return redirect('index')

    if request.method == 'POST':
        form = FormClass(request.POST)
        if form.is_valid():
            saved_obj = form.save(commit=False)
            if hasattr(saved_obj, 'created_by') and not saved_obj.created_by:
                saved_obj.created_by = request.user
            saved_obj.save()
            form.save_m2m()

            if model_name == 'customer' and request.POST.get('save_and_sync_crm'):
                from apps.crm.services import sync_master_to_client
                try:
                    client, synced = sync_master_to_client(saved_obj, user=request.user)
                    if synced and client:
                        messages.success(request, f"Data Customer berhasil disimpan dan ditransfer ke CRM / Sales ({client.client_id} - {client.company_name}).")
                    else:
                        messages.success(request, "Data Customer berhasil disimpan.")
                except Exception as e:
                    messages.warning(request, f"Data Customer tersimpan, namun sinkronisasi CRM gagal: {str(e)}")
            else:
                messages.success(request, f"Data berhasil ditambahkan.")
            return redirect(redirect_url)
    else:
        form = FormClass()
        
    template_name = get_master_template(model_name, 'form')
        
    return render(request, template_name, {
        'form': form,
        'model_name': model_name,
        'title': f"Tambah {model_name.capitalize()}"
    })

@login_required
def master_update(request, model_name, pk):
    Model, FormClass, redirect_url = get_model_and_form(model_name)
    if not Model or not FormClass:
        messages.error(request, f"Form untuk {model_name} belum tersedia.")
        return redirect('index')

    from django.shortcuts import get_object_or_404
    obj = get_object_or_404(Model, pk=pk)

    if request.method == 'POST':
        form = FormClass(request.POST, instance=obj)
        if form.is_valid():
            saved_obj = form.save(commit=False)
            if hasattr(saved_obj, 'created_by') and not saved_obj.created_by:
                saved_obj.created_by = request.user
            saved_obj.save()
            form.save_m2m()

            if model_name == 'customer' and request.POST.get('save_and_sync_crm'):
                from apps.crm.services import sync_master_to_client
                try:
                    client, synced = sync_master_to_client(saved_obj, user=request.user)
                    if synced and client:
                        messages.success(request, f"Data Customer berhasil diperbarui dan ditransfer ke CRM / Sales ({client.client_id} - {client.company_name}).")
                    else:
                        messages.success(request, "Data Customer berhasil diperbarui.")
                except Exception as e:
                    messages.warning(request, f"Data Customer diperbarui, namun sinkronisasi CRM gagal: {str(e)}")
            else:
                messages.success(request, f"Data berhasil diperbarui.")
            return redirect(redirect_url)
    else:
        form = FormClass(instance=obj)
        
    template_name = get_master_template(model_name, 'edit')
        
    return render(request, template_name, {
        'form': form,
        'obj': obj,
        'customer': obj if model_name == 'customer' else None,
        'model_name': model_name,
        'title': f"Edit {model_name.capitalize()}"
    })

@login_required
def customer_sync_to_crm(request, pk):
    from django.shortcuts import get_object_or_404
    from .models import Customer
    customer = get_object_or_404(Customer, pk=pk)
    
    from apps.crm.services import sync_master_to_client
    try:
        client, synced = sync_master_to_client(customer, user=request.user)
        if synced and client:
            messages.success(request, f'Data berhasil disinkronkan ke modul CRM / Sales ({client.client_id} - {client.company_name}).')
        else:
            messages.warning(request, 'Customer ini belum memiliki keterkaitan Client di CRM.')
    except Exception as e:
        messages.error(request, f'Gagal melakukan sinkronisasi ke CRM: {str(e)}')
        
    next_url = request.POST.get('next') or request.META.get('HTTP_REFERER')
    if next_url:
        return redirect(next_url)
    return redirect('master:master_update', model_name='customer', pk=pk)

@login_required
def master_delete(request, model_name, pk):
    Model, _, redirect_url = get_model_and_form(model_name)
    if not Model:
        return redirect('index')

    from django.shortcuts import get_object_or_404
    obj = get_object_or_404(Model, pk=pk)
    
    if request.method == 'POST':
        obj.delete()
        messages.success(request, f"Data berhasil dihapus permanen.")
        return redirect(redirect_url)

    return render(request, get_master_template(model_name, 'delete'), {
        'object': obj,
        'model_name': model_name,
        'cancel_url': redirect_url
    })

@login_required
def master_detail(request, model_name, pk):
    Model, _, redirect_url = get_model_and_form(model_name)
    if not Model:
        return redirect('index')

    from django.shortcuts import get_object_or_404
    obj = get_object_or_404(Model, pk=pk)
    
    # Simple generic detail view logic: get all fields and values
    fields = []
    for field in obj._meta.fields:
        if field.name not in ['id']:
            val = getattr(obj, field.name)
            # handle choices
            if hasattr(obj, f"get_{field.name}_display") and callable(getattr(obj, f"get_{field.name}_display")):
                val = getattr(obj, f"get_{field.name}_display")()
            elif field.name == 'branch' and val:
                from apps.master.display import clean_branch_label
                val = clean_branch_label(val)
            elif field.name in ['created_at', 'updated_at'] and val:
                val = val.strftime('%d/%m/%Y %H:%M:%S')
            elif field.name == 'created_by' and val:
                val = val.get_full_name() or val.username
            fields.append({
                'label': field.verbose_name.title() if hasattr(field, 'verbose_name') else field.name.title(),
                'value': val if val is not None and val != '' else '-'
            })
            
    return render(request, get_master_template(model_name, 'detail'), {
        'object': obj,
        'fields': fields,
        'model_name': model_name,
        'list_url': redirect_url,
        'title': f"Detail {model_name.capitalize()}"
    })


# ===== Region API Views (Cascading Dropdown Wilayah Indonesia) =====
from django.http import JsonResponse

def api_provinces(request):
    """Return all provinces for dropdown"""
    from .models import Province
    provinces = Province.objects.all().order_by('name')
    data = [{'id': p.id, 'code': p.code, 'name': p.name} for p in provinces]
    return JsonResponse({'results': data})

def api_regencies(request, province_id):
    """Return regencies filtered by province_id"""
    from .models import Regency
    regencies = Regency.objects.filter(province_id=province_id).order_by('name')
    data = [{'id': r.id, 'code': r.code, 'name': f"{r.get_type_display()} {r.name}" if r.type else r.name} for r in regencies]
    return JsonResponse({'results': data})

def api_districts(request, regency_id):
    """Return districts filtered by regency_id"""
    from .models import District
    districts = District.objects.filter(regency_id=regency_id).order_by('name')
    data = [{'id': d.id, 'code': d.code, 'name': d.name} for d in districts]
    return JsonResponse({'results': data})

def api_villages(request, district_id):
    """Return villages filtered by district_id"""
    from .models import Village
    villages = Village.objects.filter(district_id=district_id).order_by('name')
    data = [{'id': v.id, 'code': v.code, 'name': v.name, 'type': v.get_type_display()} for v in villages]
    return JsonResponse({'results': data})

def api_cities_by_province(request):
    """Return unique cities filtered by province from Coverage"""
    province = request.GET.get('province', '').strip()
    from .models import Coverage
    from apps.master.display import coverage_city_label
    qs = Coverage.objects.filter(is_active=True)
    if province:
        qs = qs.filter(province__iexact=province)
        
    unique_cities = {}
    for c in qs.order_by('city'):
        if c.city not in unique_cities:
            label = coverage_city_label(c)
            unique_cities[c.city] = {'name': c.city, 'label': label, 'tlc': c.tlc or ''}
            
    return JsonResponse({'results': list(unique_cities.values())})

def api_districts_by_city(request):
    """Return districts filtered by city name and optionally province from Coverage"""
    city = request.GET.get('city', '').strip()
    province = request.GET.get('province', '').strip()
    if not city and not province:
        return JsonResponse({'results': []})
        
    from .models import Coverage
    from apps.master.display import coverage_district_label
    qs = Coverage.objects.filter(is_active=True)
    if city:
        import re
        clean_city = re.sub(r'^(Kota|Kabupaten|Kab\.)\s+', '', city, flags=re.IGNORECASE)
        from django.db.models import Q
        qs = qs.filter(Q(city__iexact=city) | Q(city__icontains=clean_city))
    if province:
        qs = qs.filter(province__iexact=province)
        
    unique_districts = {}
    for c in qs.order_by('district'):
        if c.district not in unique_districts:
            label = coverage_district_label(c)
            tlc_with_id = f"{c.tlc}{c.id}" if c.tlc else str(c.id)
            unique_districts[c.district] = {'id': c.id, 'name': c.district, 'label': label, 'tlc': tlc_with_id}
            
    return JsonResponse({'results': list(unique_districts.values())})


@login_required
def price_import(request):
    """
    Import Harga Publish (Publish Rates) from Excel template.
    Supports Sheet1 (Rates) and Sheet2 (Coverage Reference).
    """
    from .models import Coverage, Service, Price
    from apps.master.display import coverage_city_label
    import openpyxl
    from django.db import transaction
    from django.contrib import messages
    import re

    coverages_list = Coverage.objects.filter(is_active=True).order_by('city')
    # Unique cities / key hub coverages for default origin selection
    hub_coverages = []
    seen_cities = set()
    for c in coverages_list:
        clean_c = re.sub(r'(?i)^(kota|kabupaten)\s+', '', c.city).strip()
        if clean_c not in seen_cities:
            seen_cities.add(clean_c)
            hub_coverages.append({'id': c.id, 'label': coverage_city_label(c), 'city': c.city})

    if request.method == 'POST':
        excel_file = request.FILES.get('excel_file')
        if not excel_file:
            messages.error(request, "Silakan pilih file Excel (.xlsx) untuk diimpor.")
            return redirect('master:price_import')

        default_origin_id = request.POST.get('default_origin_id')
        auto_create_coverage = request.POST.get('auto_create_coverage') == '1'
        overwrite_existing = request.POST.get('overwrite_existing') == '1'

        default_origin = None
        if default_origin_id:
            default_origin = Coverage.objects.filter(pk=default_origin_id, is_active=True).first()
        if not default_origin:
            default_origin = Coverage.objects.filter(is_active=True, city__icontains='Jakarta').first() or Coverage.objects.filter(is_active=True).first()

        try:
            wb = openpyxl.load_workbook(excel_file, data_only=True)
        except Exception as e:
            messages.error(request, f"Gagal membaca file Excel: {str(e)}")
            return redirect('master:price_import')

        # 1. Preload / Build Coverage Lookups
        all_covs = list(Coverage.objects.filter(is_active=True))
        cov_by_city_dist = {}
        cov_by_dist = {}
        cov_by_city = {}
        cov_by_tlc = {}

        for c in all_covs:
            clean_c = re.sub(r'(?i)^(kota|kabupaten)\s+', '', c.city).strip().upper()
            clean_d = c.district.strip().upper()
            cov_by_city_dist[(clean_c, clean_d)] = c
            if clean_d not in cov_by_dist:
                cov_by_dist[clean_d] = c
            if clean_c not in cov_by_city:
                cov_by_city[clean_c] = c
            if c.tlc:
                cov_by_tlc[c.tlc.strip().upper()] = c

        # 2. Parse Sheet2 (Wilayah Reference) if present to supplement Coverages
        sheet2_dist_lookup = {}
        if 'Sheet2' in wb.sheetnames:
            ws2 = wb['Sheet2']
            headers2 = [str(ws2.cell(1, col).value or '').strip() for col in range(1, ws2.max_column + 1)]
            
            # Map column names
            h2_map = {name.lower(): idx + 1 for idx, name in enumerate(headers2)}
            col_concate = h2_map.get('concate city district', 1)
            col_tlc = h2_map.get('tlc')
            col_dist_code = h2_map.get('district_code')
            col_city_code = h2_map.get('city_code')
            col_kec = h2_map.get('sub_district_name(kecamatan)')
            col_city = h2_map.get('city_name')
            col_prov = h2_map.get('province_name')
            col_pos = h2_map.get('postal_code')

            new_coverages_to_create = []
            seen_new_covs = set()
            for r in range(2, ws2.max_row + 1):
                kec_val = str(ws2.cell(r, col_kec).value or '').strip() if col_kec else ''
                city_val = str(ws2.cell(r, col_city).value or '').strip() if col_city else ''
                prov_val = str(ws2.cell(r, col_prov).value or '').strip() if col_prov else ''
                pos_val = str(ws2.cell(r, col_pos).value or '').strip() if col_pos else ''
                tlc_val = str(ws2.cell(r, col_tlc).value or '').strip() if col_tlc else ''

                clean_c = re.sub(r'(?i)^(kota|kabupaten)\s+', '', city_val).strip().upper()
                clean_d = kec_val.strip().upper()
                matched_c = cov_by_city_dist.get((clean_c, clean_d)) or cov_by_dist.get(clean_d)
                
                if not matched_c and auto_create_coverage and clean_d and clean_c:
                    pair_key = (clean_c, clean_d)
                    if pair_key not in seen_new_covs:
                        seen_new_covs.add(pair_key)
                        new_coverages_to_create.append(Coverage(
                            district=kec_val or clean_d,
                            city=city_val or clean_c,
                            province=prov_val or 'Indonesia',
                            postal_code=pos_val or None,
                            tlc=tlc_val or None,
                            is_covered=True,
                            is_verified=True,
                            is_active=True
                        ))

            if new_coverages_to_create:
                Coverage.objects.bulk_create(new_coverages_to_create, ignore_conflicts=True)

            # Reload all coverages from DB with saved PKs
            all_covs = list(Coverage.objects.filter(is_active=True))
            cov_by_city_dist.clear()
            cov_by_dist.clear()
            cov_by_city.clear()
            cov_by_tlc.clear()
            for c in all_covs:
                clean_c = re.sub(r'(?i)^(kota|kabupaten)\s+', '', c.city).strip().upper()
                clean_d = c.district.strip().upper()
                cov_by_city_dist[(clean_c, clean_d)] = c
                if clean_d not in cov_by_dist:
                    cov_by_dist[clean_d] = c
                if clean_c not in cov_by_city:
                    cov_by_city[clean_c] = c
                if c.tlc:
                    cov_by_tlc[c.tlc.strip().upper()] = c

            # Build sheet2_dist_lookup directly with saved Coverage objects
            for r in range(2, ws2.max_row + 1):
                dist_code_val = str(ws2.cell(r, col_dist_code).value or '').strip() if col_dist_code else ''
                city_code_val = str(ws2.cell(r, col_city_code).value or '').strip() if col_city_code else ''
                kec_val = str(ws2.cell(r, col_kec).value or '').strip() if col_kec else ''
                city_val = str(ws2.cell(r, col_city).value or '').strip() if col_city else ''
                clean_c = re.sub(r'(?i)^(kota|kabupaten)\s+', '', city_val).strip().upper()
                clean_d = kec_val.strip().upper()
                matched_c = cov_by_city_dist.get((clean_c, clean_d)) or cov_by_dist.get(clean_d) or cov_by_city.get(clean_c)
                if matched_c and matched_c.pk:
                    if dist_code_val:
                        sheet2_dist_lookup[dist_code_val.upper()] = matched_c
                    if city_code_val:
                        sheet2_dist_lookup[city_code_val.upper()] = matched_c

        # 3. Preload Services
        all_services = {s.code.upper(): s for s in Service.objects.all()}
        service_aliases = {
            'UDRREG': 'REG',
            'UDRODS': 'ODS',
            'UDRSDS': 'SDS',
            'DARAT': 'RD',
            'LAUT': 'RL',
        }

        # 4. Process Sheet1 (Rates)
        sheet1_name = 'Sheet1' if 'Sheet1' in wb.sheetnames else wb.sheetnames[0]
        ws1 = wb[sheet1_name]

        headers1 = [str(ws1.cell(1, col).value or '').strip().lower() for col in range(1, ws1.max_column + 1)]
        h1_map = {name: idx + 1 for idx, name in enumerate(headers1)}
        
        col_svc = h1_map.get('service_code', 3)
        col_orig = h1_map.get('origin', 4)
        col_dest = h1_map.get('destination_name', 6)
        col_dest_code = h1_map.get('destination_district_code', 7)
        col_sla_min = h1_map.get('sla_min', 8)
        col_sla_max = h1_map.get('sla_max', 9)
        col_min_w = h1_map.get('min_weight', 10)
        col_p_kg = h1_map.get('price_kg', 12)

        # Preload existing Prices to check overwrite / duplicate
        existing_prices = {}
        for p in Price.objects.all().select_related('origin', 'destination', 'service'):
            k = (p.origin_id, p.destination_id, p.service_id)
            existing_prices[k] = p

        created_count = 0
        updated_count = 0
        skipped_count = 0
        prices_to_create = []
        prices_to_update = []

        for r in range(2, ws1.max_row + 1):
            p_kg_val = ws1.cell(r, col_p_kg).value
            if p_kg_val is None or str(p_kg_val).strip() == '':
                skipped_count += 1
                continue
            try:
                price_per_kg = float(p_kg_val)
            except (ValueError, TypeError):
                skipped_count += 1
                continue

            # Resolve Service
            svc_code_raw = str(ws1.cell(r, col_svc).value or 'REG').strip().upper()
            target_svc_code = service_aliases.get(svc_code_raw, svc_code_raw)
            service_obj = all_services.get(target_svc_code) or all_services.get(svc_code_raw)
            if not service_obj:
                service_obj = Service.objects.create(
                    code=target_svc_code,
                    name=f"Layanan {target_svc_code}",
                    is_active=True,
                    is_verified=True
                )
                all_services[target_svc_code] = service_obj

            # Resolve Destination Coverage
            dest_name_val = str(ws1.cell(r, col_dest).value or '').strip()
            dest_code_val = str(ws1.cell(r, col_dest_code).value or '').strip() if col_dest_code else ''
            dest_cov = None

            if dest_code_val and dest_code_val.upper() in sheet2_dist_lookup:
                lookup_res = sheet2_dist_lookup[dest_code_val.upper()]
                if isinstance(lookup_res, Coverage):
                    dest_cov = lookup_res
                elif isinstance(lookup_res, tuple):
                    dest_cov = cov_by_city_dist.get(lookup_res)

            if not dest_cov and dest_name_val:
                parts = [p.strip().upper() for p in dest_name_val.split(',')]
                if len(parts) >= 2:
                    c_clean = re.sub(r'(?i)^(kota|kabupaten)\s+', '', parts[0]).strip().upper()
                    d_clean = parts[1].strip().upper()
                    dest_cov = cov_by_city_dist.get((c_clean, d_clean)) or cov_by_dist.get(d_clean) or cov_by_city.get(c_clean)
                elif len(parts) == 1:
                    d_clean = parts[0].strip().upper()
                    dest_cov = cov_by_dist.get(d_clean) or cov_by_city.get(d_clean)

            if not dest_cov:
                skipped_count += 1
                continue

            # Resolve Origin Coverage
            orig_val = str(ws1.cell(r, col_orig).value or '').strip()
            orig_cov = None
            if orig_val and orig_val.lower() not in ['null', 'none', 'all', '']:
                orig_cov = sheet2_dist_lookup.get(orig_val.upper()) if orig_val.upper() in sheet2_dist_lookup else None
                if isinstance(orig_cov, tuple):
                    orig_cov = cov_by_city_dist.get(orig_cov)
                if not orig_cov:
                    orig_cov = cov_by_dist.get(orig_val.upper()) or cov_by_city.get(orig_val.upper())
            
            if not orig_cov:
                orig_cov = default_origin

            if not orig_cov:
                skipped_count += 1
                continue

            # Min Weight & SLA
            min_w_val = ws1.cell(r, col_min_w).value
            min_weight = float(min_w_val) if min_w_val is not None else 1.0

            sla_min_val = ws1.cell(r, col_sla_min).value
            sla_max_val = ws1.cell(r, col_sla_max).value
            if sla_min_val is not None and sla_max_val is not None:
                if str(sla_min_val) == str(sla_max_val):
                    est_days = f"{sla_min_val} Hari"
                else:
                    est_days = f"{sla_min_val}-{sla_max_val} Hari"
            elif sla_min_val is not None:
                est_days = f"{sla_min_val} Hari"
            else:
                est_days = "2-3 Hari"

            key = (orig_cov.pk, dest_cov.pk, service_obj.pk)
            existing_price = existing_prices.get(key)

            if existing_price:
                if overwrite_existing:
                    existing_price.price_per_kg = price_per_kg
                    existing_price.min_weight = min_weight
                    existing_price.estimated_days = est_days
                    existing_price.is_active = True
                    prices_to_update.append(existing_price)
                    updated_count += 1
                else:
                    skipped_count += 1
            else:
                new_p = Price(
                    origin=orig_cov,
                    destination=dest_cov,
                    service=service_obj,
                    price_per_kg=price_per_kg,
                    min_weight=min_weight,
                    estimated_days=est_days,
                    is_active=True
                )
                prices_to_create.append(new_p)
                existing_prices[key] = new_p
                created_count += 1

        # Execute bulk save in transactions
        with transaction.atomic():
            if prices_to_create:
                chunk_size = 1000
                for i in range(0, len(prices_to_create), chunk_size):
                    Price.objects.bulk_create(prices_to_create[i:i + chunk_size])
            if prices_to_update:
                chunk_size = 1000
                for i in range(0, len(prices_to_update), chunk_size):
                    Price.objects.bulk_update(prices_to_update[i:i + chunk_size], ['price_per_kg', 'min_weight', 'estimated_days', 'is_active'])

        messages.success(
            request,
            f"Proses impor berhasil selesai! {created_count:,} tarif baru ditambahkan, "
            f"{updated_count:,} diperbarui, dan {skipped_count:,} baris dilewati."
        )
        return redirect('master:price')

    return render(request, 'master/prices/price_import.html', {
        'hub_coverages': hub_coverages[:60],
        'title': 'Import Harga Publish (Publish Rate)'
    })


@login_required
def price_download_template(request):
    """Download base reference template for Publish Rate import."""
    import os
    from django.conf import settings
    from django.http import HttpResponse, Http404

    template_path = os.path.join(settings.BASE_DIR, 'upload_template_harga_publish (1).xlsx')
    if not os.path.exists(template_path):
        raise Http404("File template tidak ditemukan.")

    with open(template_path, 'rb') as f:
        file_data = f.read()

    response = HttpResponse(file_data, content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename="Template_Import_Harga_Publish_Paketin.xlsx"'
    return response